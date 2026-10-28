"""AX 전문컨설턴트 3급 연습 문제 10개(문제지)와 정답·해설 파일 만들기. 실행: python3 make_exam.py"""
from docx import Document
from docx.shared import Pt,Mm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

MC=[ # (문제, [보기4], 정답번호, 해설)
("AX(AI Transformation, 인공지능 전환)에 대한 설명으로 가장 알맞은 것은?",
 ["종이 문서를 컴퓨터 파일로 바꾸는 일","인공지능을 업무·서비스·조직 전반에 도입해 일하는 방식과 가치를 바꾸는 일","컴퓨터 한 대에 새 프로그램을 설치하는 일","직원 수를 줄이는 일"],2,
 "AX는 인공지능을 일부 업무에 써 보는 데서 그치지 않고, 업무·서비스·조직이 일하는 방식 자체를 바꾸어 새로운 가치를 만드는 것을 말해요. ‘종이 문서를 컴퓨터 파일로 바꾸는 일’은 디지털화에 가까운 설명이고, ‘직원 수를 줄이는 일’은 AX의 목적이 아니에요."),
("생성형 AI에 대한 설명으로 알맞은 것은?",
 ["데이터를 저장하는 장치","정해진 규칙만 되풀이해서 실행하는 프로그램","글, 이미지, 코드처럼 새로운 결과물을 만들어 내는 인공지능","인터넷 속도를 높이는 기술"],3,
 "생성형 AI는 사람이 준 질문이나 지시(프롬프트)를 바탕으로 글, 이미지, 코드 등 새로운 결과물을 만들어 내는 AI예요. ChatGPT, Gemini 같은 서비스가 대표적이에요."),
("머신러닝(기계학습)을 가장 잘 설명한 것은?",
 ["사람이 모든 규칙을 일일이 프로그래밍하는 방식","컴퓨터가 많은 데이터를 보고 스스로 규칙과 패턴을 배우는 방식","기계를 사람이 직접 조작하는 방식","컴퓨터를 오래 켜 두어도 고장 나지 않게 하는 방식"],2,
 "머신러닝은 데이터에서 규칙과 패턴을 스스로 학습해서 예측이나 판단을 하는 기술이에요. ‘사람이 모든 규칙을 일일이 프로그래밍하는 방식’은 전통적인 프로그래밍이에요."),
("생성형 AI에게 원하는 일을 시키려고 입력하는 질문이나 지시문을 무엇이라고 하는가?",
 ["프로세서","프롬프트","프로토콜","프로그램"],2,
 "프롬프트(prompt)는 AI에게 하는 질문이나 지시문이에요. 역할, 목표, 재료, 규칙을 구체적으로 쓸수록 원하는 결과를 얻기 쉬워요."),
("회사 자료를 외부의 무료 AI 서비스에 입력해야 할 때 가장 알맞은 행동은?",
 ["거래처 이름과 금액을 그대로 모두 입력한다","이름·금액 같은 민감한 정보는 가리거나 빼고 입력한다","비밀번호를 함께 입력해서 AI가 파일을 열게 한다","누가 올렸는지 알 수 없으니 신경 쓰지 않는다"],2,
 "외부 AI에 입력한 내용은 회사 밖으로 나가요. 거래처명, 개인정보, 영업 비밀 같은 민감한 정보는 가리거나(마스킹) 빼고 쓰는 것이 기본이에요. 회사의 정보 보안 규정이 있으면 먼저 확인해야 해요."),
("기업이 AX를 추진하는 순서로 가장 알맞은 것은?",
 ["확산·정착 → 시범 적용 → 과제 발굴 → 현황 진단","시범 적용 → 현황 진단 → 확산·정착 → 과제 발굴","현황 진단 → 과제 발굴(우선순위 정하기) → 시범 적용 → 확산·정착","과제 발굴 → 확산·정착 → 현황 진단 → 시범 적용"],3,
 "먼저 현재 업무와 데이터 상태를 진단하고, AI로 풀 과제를 찾아 우선순위를 정한 다음, 작게 시범 적용(PoC)해 보고, 효과가 확인되면 조직 전체로 확산·정착시키는 것이 일반적인 순서예요."),
("AI에게 엑셀 자료를 분석시키기 전에 가장 먼저 해야 할 일로 알맞은 것은?",
 ["그래프 색을 예쁘게 고른다","자료의 중복, 빈칸, 표기 흔들림(예: 서울/서울시)을 정리한다","파일 이름을 영어로 바꾼다","자료를 최대한 많이 늘려서 넣는다"],2,
 "데이터가 지저분하면 AI도 틀린 결과를 내기 쉬워요. 중복을 지우고, 날짜·숫자 모양을 통일하고, 같은 뜻의 값을 하나로 맞추는 '데이터 정리'가 분석의 첫 단계예요."),
("AI 도입의 효과를 확인하는 방법으로 가장 알맞은 것은?",
 ["직원들이 좋아하는 것 같다고 추측한다","도입 전과 후의 업무 시간, 오류 건수 같은 숫자를 비교한다","AI 서비스를 쓰는 사람 수가 늘었다는 것만 본다","다른 회사가 도입했다는 사실만 확인한다"],2,
 "효과는 '측정할 수 있는 지표'로 확인해야 해요. 예를 들어 월 마감에 걸리던 시간이 8시간에서 3시간으로 줄었는지, 입력 오류가 몇 건 줄었는지처럼 도입 전후를 숫자로 비교해요."),
]
# 정답 번호가 한 곳에 몰리지 않게 보기 순서를 섞음 (정답 위치를 아래 목록대로 맞춤)
WANT=[2,3,1,4,3,4,1,2]
for i,w in enumerate(WANT):
    q,opts,ans,exp=MC[i];opts=opts[:];a,b=ans-1,w-1;opts[a],opts[b]=opts[b],opts[a];MC[i]=(q,opts,w,exp)
SA=[ # (문제, 유형, 모범답안, 채점기준)
("생성형 AI가 사실이 아닌 내용을 사실처럼 그럴듯하게 만들어 내는 현상을 무엇이라고 하는지 쓰시오.","단답형",
 "환각 (할루시네이션, hallucination)",
 "‘환각’ 또는 ‘할루시네이션/hallucination’이면 10점. 비슷한 설명만 쓰고 용어가 없으면 5점. 철자가 조금 틀려도 의미가 같으면 인정."),
("회사에서 생성형 AI를 사용할 때 지켜야 할 주의사항을 두 가지 쓰시오.","서술형",
 "① 개인정보·거래처명·영업 비밀 같은 민감한 정보는 입력하지 않거나 가려서 입력한다. ② AI의 답은 틀릴 수 있으므로 숫자와 사실을 사람이 직접 확인한다. (그 밖에 저작권·출처 확인, 회사 보안 규정 준수 등)",
 "알맞은 주의사항 1가지당 5점, 최대 10점. 가능한 답: 민감정보 입력 금지/가리기, 결과의 사실·숫자 확인(검증), 저작권·출처 확인, 회사 보안 규정 준수, 최종 책임은 사람에게 있음."),
]
def font(run,size=11,bold=False,color=None):
    run.font.name='맑은 고딕';run._element.rPr.rFonts.set(qn('w:eastAsia'),'맑은 고딕');run.font.size=Pt(size);run.bold=bold
    if color:run.font.color.rgb=RGBColor(*color)
def para(doc,text='',size=11,bold=False,align=None,space=4,color=None,indent=None):
    p=doc.add_paragraph();r=p.add_run(text);font(r,size,bold,color)
    p.paragraph_format.space_after=Pt(space);p.paragraph_format.space_before=Pt(0)
    if align:p.alignment=align
    if indent is not None:p.paragraph_format.left_indent=Mm(indent)
    return p
def newdoc():
    d=Document();s=d.sections[0];s.page_width=Mm(210);s.page_height=Mm(297)
    for a in('left_margin','right_margin'):setattr(s,a,Mm(20))
    s.top_margin=Mm(18);s.bottom_margin=Mm(18)
    st=d.styles['Normal'];st.font.name='맑은 고딕';st.element.rPr.rFonts.set(qn('w:eastAsia'),'맑은 고딕');st.font.size=Pt(11)
    return d
CIR='①②③④'
TITLE='인공지능전환(AX) 전문컨설턴트 3급 연습 문제'
# ---- 문제지
d=newdoc()
para(d,TITLE,18,True,WD_ALIGN_PARAGRAPH.CENTER,6)
para(d,'(문제지)',12,False,WD_ALIGN_PARAGRAPH.CENTER,10,(90,90,90))
para(d,'응시자 이름: ____________________      점수: ________ / 100점',11,False,None,4)
para(d,'시험 시간: 20분   ·   객관식 8문항(문항당 10점, 80점) + 주관식 2문항(문항당 10점, 20점)',10.5,False,None,12,(90,90,90))
para(d,'[1~8] 알맞은 것을 하나만 고르시오. (각 10점)',12,True,None,8)
for i,(q,opts,ans,exp) in enumerate(MC,1):
    para(d,f'{i}. {q}',11.5,True,None,3)
    for j,o in enumerate(opts):para(d,f'{CIR[j]} {o}',11,False,None,1,None,6)
    para(d,'',6,False,None,4)
para(d,'[9~10] 알맞은 답을 쓰시오. (각 10점)',12,True,None,8)
for k,(q,t,ans,rub) in enumerate(SA,9):
    para(d,f'{k}. {q}  ({t})',11.5,True,None,6)
    n=2 if t=='단답형' else 5
    for _ in range(n):para(d,'_'*70,11,False,None,6,(150,150,150))
    para(d,'',6,False,None,6)
d.save(f'AX전문컨설턴트3급_연습문제_문제지.docx')
# ---- 정답·해설
d=newdoc()
para(d,TITLE,18,True,WD_ALIGN_PARAGRAPH.CENTER,6)
para(d,'(정답 및 해설 · 강사용)',12,False,WD_ALIGN_PARAGRAPH.CENTER,10,(90,90,90))
para(d,'1. 정답표',13,True,None,6)
t=d.add_table(rows=2,cols=9);t.style='Table Grid'
for i in range(8):
    t.rows[0].cells[i+1].text=f'{i+1}번';t.rows[1].cells[i+1].text=CIR[MC[i][2]-1]
t.rows[0].cells[0].text='문항';t.rows[1].cells[0].text='정답'
for row in t.rows:
    for c in row.cells:
        for p in c.paragraphs:
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:font(r,10.5,True)
para(d,'',6,False,None,4)
para(d,'주관식 9~10번은 아래 모범 답안과 채점 기준을 참고하세요.',10.5,False,None,10,(90,90,90))
para(d,'2. 객관식 해설',13,True,None,6)
for i,(q,opts,ans,exp) in enumerate(MC,1):
    para(d,f'{i}번  정답 {CIR[ans-1]}  {opts[ans-1]}',11.5,True,None,2)
    para(d,exp,11,False,None,8,None,4)
para(d,'3. 주관식 모범 답안과 채점 기준',13,True,None,6)
for k,(q,tp,ans,rub) in enumerate(SA,9):
    para(d,f'{k}번 ({tp}) {q}',11.5,True,None,2)
    para(d,'모범 답안: '+ans,11,False,None,2,None,4)
    para(d,'채점 기준: '+rub,11,False,None,8,(90,90,90),4)
para(d,'4. 출제 안내',13,True,None,6)
for x in ['난이도: 하(기본 개념과 업무 상식 수준). 합격선을 정할 때는 60점 이상을 기준으로 보시길 권해요.','이 문제는 새로 만든 연습 문제이며, 실제 자격시험 문제나 공식 출제 범위와는 다를 수 있어요. 시험을 시행하는 기관의 출제 기준을 먼저 확인해 주세요.','한글(HWP)로 쓰실 때는 이 문서를 열어 내용을 복사해서 붙여 넣으면 돼요.']:
    para(d,'· '+x,10.5,False,None,3,(70,70,70))
d.save('AX전문컨설턴트3급_연습문제_정답및해설.docx')
print('ok')
